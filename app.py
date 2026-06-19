"""
Flask Web Server cho LangGraph Chatbot
"""
import os
import time
import json
import uuid
import shutil
import importlib
from datetime import datetime
from pathlib import Path

from chatbot import PROVIDER
from flask import Flask, request, jsonify, render_template
from flask_cors import CORS
from dotenv import load_dotenv
from werkzeug.utils import secure_filename
from backend import ChatHistoryStore, ProcessInspector, UserDataStore

# Load environment
load_dotenv()
ENABLE_OCR = os.getenv("ENABLE_OCR", "false").strip().lower() in ("1", "true", "yes")

app = Flask(__name__, template_folder="templates", static_folder="static")
CORS(app)
history_store = ChatHistoryStore()
user_store = UserDataStore()
process_inspector = ProcessInspector(history_store)
# ══════════════════════════════════════════════════════════════════
#  LAZY IMPORT & INITIALIZATION
# ══════════════════════════════════════════════════════════════════
_graph = None
_db = None
_provider = None
_chatbot_module = None


def init_chatbot(
    provider_override: str | None = None,
) -> None:
    """Initialize chatbot on first request (lazy loading)."""
    global _graph, _db, _provider, _chatbot_module
    if _graph is not None and provider_override is None:
        return

    # If we already have a graph and the override matches current provider, keep existing graph
    if (
        _graph is not None
        and provider_override is not None
        and provider_override == _provider
    ):
        return

    # Rebuild graph only when provider changes
    if (
        _graph is not None
        and provider_override is not None
        and provider_override != _provider
    ):
        _graph = None
        _db = None
        _chatbot_module = None

    # Import the chatbot module
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "chatbot_module",
        os.path.join(os.path.dirname(__file__), "multi_agent_chatbot_v3.py")
    )
    chatbot_module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(chatbot_module)

    if provider_override is not None:
        if provider_override not in {"groq", "anthropic"}:
            raise RuntimeError(f"Unsupported provider: {provider_override}")
        if (
            provider_override == "groq"
            and not getattr(chatbot_module, "GROQ_KEY", None)
        ):
            raise RuntimeError(
                "GROQ_API_KEY not configured for GROQ provider."
            )
        if (
            provider_override == "anthropic"
            and not getattr(chatbot_module, "ANTHROPIC_KEY", None)
        ):
            raise RuntimeError(
                "ANTHROPIC_API_KEY not configured for Anthropic provider."
            )
        chatbot_module.PROVIDER = provider_override
        chatbot_module.USE_GROQ = provider_override == "groq"

    _graph = chatbot_module.build_graph()
    _db = None
    chatbot_module.db = None
    _chatbot_module = chatbot_module

    # Prefer provider/key info from the chatbot module (module owns config)
    _check_gemini_packages = getattr(chatbot_module, "_check_gemini_packages", lambda: None)
    _provider = getattr(chatbot_module, "PROVIDER", PROVIDER)

    if _provider == "none":
        raise RuntimeError("No API Key configured. Check .env file.")

    print(f"[INFO] Initializing chatbot with provider={_provider}")
    try:
        _check_gemini_packages()
    except Exception:
        # Non-fatal: package checks are informational only
        pass
    print("[INFO] LangGraph compiled successfully")


def handle_query(session_id: str, query: str, provider_override: str | None = None, user_id: str | None = None) -> dict:
    """
    Run a single query through the chatbot and return the result.
    
    Returns:
        {
            "session_id": str,
            "final_answer": str,
            "intent": str,
            "agent_timings": List[dict],
            "log": List[str],
        }
    """
    if _graph is None or provider_override is not None:
        init_chatbot(provider_override)

    from multi_agent_chatbot_v3 import ChatState, ChatPersistence

    global _db
    if _db is None or getattr(_db, "session_id", None) != session_id:
        _db = ChatPersistence(session_id)
        if _chatbot_module is not None:
            _chatbot_module.db = _db

    # Create initial state
    turn_id = int(time.time() * 1000) % 100000
    initial: ChatState = {
        "session_id": session_id,
        "turn_id": turn_id,
        "user_query": query,
        "user_id": user_id or "",
        "intent": "",
        "goal": "",
        "kb_context": "",
        "workspace_messages": [],
        "coordination_summary": "",
        "planning": "",
        "validation": "",
        "web_results": "",
        "calc_result": "",
        "datetime_result": "",
        "analysis": "",
        "final_answer": "",
        "conversation_history": [],
        "agent_timings": [],
        "call_graph": [],
        "log": [],
    }

    # Run through graph
    try:
        t0 = time.perf_counter()
        result = _graph.invoke(initial)
        elapsed = time.perf_counter() - t0

        return {
            "session_id": session_id,
            "final_answer": result.get("final_answer", ""),
            "intent": result.get("intent", ""),
            "agent_timings": result.get("agent_timings", []),
            "log": result.get("log", []),
            "elapsed": round(elapsed, 3),
        }
    except Exception:
        import traceback
        traceback.print_exc()
        raise


def _get_uploaded_file_path(url: str | None) -> Path | None:
    if not url:
        return None
    normalized = url.strip().split("?", 1)[0].split("#", 1)[0]
    if normalized.startswith("http://") or normalized.startswith("https://"):
        if "/static/" in normalized:
            normalized = normalized.split("/static/", 1)[-1]
        else:
            normalized = "/".join(normalized.split("/")[3:])
    normalized = normalized.lstrip("/")
    if normalized.startswith("static/"):
        normalized = normalized[len("static/"):]

    static_root = Path(app.root_path) / app.static_folder
    if not static_root.is_absolute():
        static_root = (Path(__file__).resolve().parent / static_root).resolve()

    candidate = (static_root / normalized).resolve()
    if candidate.exists():
        return candidate

    alternate = (Path(__file__).resolve().parent / app.static_folder / normalized).resolve()
    if alternate.exists():
        return alternate

    print(
        f"[Attachment] file path not found: url={url}, normalized={normalized}, "
        f"candidate={candidate}, alternate={alternate}"
    )
    return None


def _truncate_text(text: str, max_chars: int = 4000) -> str:
    if not text:
        return ""
    clean = text.strip()
    if len(clean) <= max_chars:
        return clean
    return clean[:max_chars].rstrip() + "\n\n[...nội dung bị cắt bớt... ]"


def _find_tesseract_binary_path() -> str | None:
    candidate = shutil.which("tesseract")
    if candidate:
        return candidate

    env_paths = [
        os.environ.get("TESSERACT_CMD"),
        os.environ.get("TESSERACT_PATH"),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    for path in env_paths:
        if path:
            if Path(path).exists():
                return str(Path(path).resolve())
    return None


def _get_extraction_dependencies() -> dict[str, bool]:
    deps = {
        "PIL": False,
        "pytesseract": False,
        "easyocr": False,
        "tesseract_binary": False,
        "PyPDF2": False,
        "pdf2image": False,
        "PyMuPDF": False,
        "python-docx": False,
        "openpyxl": False,
        "ocr_enabled": ENABLE_OCR,
    }

    def _module_available(name: str) -> bool:
        try:
            return importlib.util.find_spec(name) is not None
        except Exception:
            return False

    deps["PIL"] = _module_available("PIL")
    deps["pytesseract"] = _module_available("pytesseract")
    deps["easyocr"] = _module_available("easyocr")
    deps["PyPDF2"] = _module_available("PyPDF2")
    deps["pdf2image"] = _module_available("pdf2image")
    deps["PyMuPDF"] = _module_available("fitz")
    deps["python-docx"] = _module_available("docx")
    deps["openpyxl"] = _module_available("openpyxl")
    deps["tesseract_binary"] = _find_tesseract_binary_path() is not None
    return deps


def _run_easyocr(image) -> str | None:
    if not ENABLE_OCR:
        return None
    try:
        from easyocr import Reader
    except Exception as exc:
        print(f"[OCR] easyocr not available: {exc}")
        return None
    try:
        reader = Reader(["vi", "en"], gpu=False)
        result = reader.readtext(image)
        if not result:
            return None
        return "\n".join([item[1] for item in result if item and item[1].strip()])
    except Exception as exc:
        print(f"[OCR] easyocr extraction error: {exc}")
        return None


def _extract_text_from_image(path: Path) -> str | None:
    if not ENABLE_OCR:
        print("[OCR] OCR disabled via ENABLE_OCR environment setting.")
        return None

    try:
        from PIL import Image, ImageOps
    except Exception as exc:
        print(f"[OCR] Missing Pillow dependency: {exc}")
        return None

    pytesseract_available = False
    try:
        import pytesseract
        pytesseract_available = True
    except Exception as exc:
        print(f"[OCR] pytesseract not installed: {exc}")

    def _ocr_with_tesseract(img):
        if not pytesseract_available:
            return None
        try:
            return pytesseract.image_to_string(img, config="--psm 6 --oem 1")
        except Exception as exc:
            print(f"[OCR] pytesseract error: {exc}")
            return None

    def _try_easyocr(img):
        return _run_easyocr(img)

    try:
        with Image.open(path) as img:
            img = img.convert("L")
            img = ImageOps.autocontrast(img)
            if img.width < 1200 or img.height < 1200:
                img = img.resize((min(img.width * 2, 1600), min(img.height * 2, 1600)), Image.LANCZOS)

            if pytesseract_available and shutil.which("tesseract"):
                text = _ocr_with_tesseract(img)
                if text and text.strip():
                    return _truncate_text(text)
                inverted = ImageOps.invert(img)
                text = _ocr_with_tesseract(inverted)
                if text and text.strip():
                    return _truncate_text(text)

            text = _try_easyocr(img)
            if text and text.strip():
                return _truncate_text(text)

            text = _try_easyocr(ImageOps.invert(img))
            if text and text.strip():
                return _truncate_text(text)

            return None
    except Exception as exc:
        import traceback
        traceback.print_exc()
        print(f"[OCR] Error extracting image text from {path}: {exc}")
        return None


def _extract_text_from_pdf(path: Path) -> str | None:
    text_pages = []
    try:
        import PyPDF2
    except Exception as exc:
        print(f"[PDF] Missing PyPDF2 dependency: {exc}")
        return None

    try:
        reader = PyPDF2.PdfReader(str(path))
        for page in reader.pages:
            try:
                extracted = page.extract_text()
                if extracted:
                    text_pages.append(extracted)
            except Exception:
                continue
    except Exception as exc:
        print(f"[PDF] PyPDF2 read error: {exc}")

    if text_pages:
        return _truncate_text("\n\n".join(text_pages))

    if not ENABLE_OCR:
        return None

    fallback_texts = []
    try:
        from pdf2image import convert_from_path
        images = convert_from_path(str(path), dpi=200)
        for img in images:
            text = _extract_text_from_pil_image(img)
            if text:
                fallback_texts.append(text)
        if fallback_texts:
            return _truncate_text("\n\n".join(fallback_texts))
    except Exception as exc:
        print(f"[PDF] pdf2image unavailable or failed: {exc}")

    try:
        import fitz
        doc = fitz.open(str(path))
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            text = _extract_text_from_pil_image(pix.get_pil_image())
            if text:
                fallback_texts.append(text)
        if fallback_texts:
            return _truncate_text("\n\n".join(fallback_texts))
    except Exception as exc:
        print(f"[PDF] PyMuPDF fallback failed: {exc}")

    return None


def _extract_text_from_pil_image(image) -> str | None:
    if not ENABLE_OCR:
        return None

    pytesseract_available = False
    try:
        import pytesseract
        pytesseract_available = True
    except Exception as exc:
        print(f"[OCR] pytesseract unavailable for PIL image OCR: {exc}")

    if pytesseract_available:
        try:
            text = pytesseract.image_to_string(image, config="--psm 6 --oem 1")
            if text and text.strip():
                return _truncate_text(text)
        except Exception as exc:
            print(f"[OCR] PIL image pytesseract error: {exc}")

    easy_text = _run_easyocr(image)
    if easy_text and easy_text.strip():
        return _truncate_text(easy_text)

    return None


def _extract_text_from_docx(path: Path) -> str | None:
    try:
        import docx
    except Exception as exc:
        print(f"[DOCX] Missing python-docx dependency: {exc}")
        return None
    try:
        document = docx.Document(str(path))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        return _truncate_text("\n\n".join(paragraphs))
    except Exception as exc:
        print(f"[DOCX] Error extracting DOCX text: {exc}")
        return None


def _extract_text_from_xlsx(path: Path) -> str | None:
    try:
        import openpyxl
    except Exception as exc:
        print(f"[XLSX] Missing openpyxl dependency: {exc}")
        return None
    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            rows.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) for cell in row if cell is not None]
                if values:
                    rows.append("\t".join(values))
        return _truncate_text("\n".join(rows))
    except Exception as exc:
        print(f"[XLSX] Error reading workbook {path}: {exc}")
        return None


def _extract_text_from_file(path: Path, attachment_type: str | None = None) -> str | None:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".ts", ".java", ".yaml", ".yml", ".html", ".htm"}:
        try:
            return _truncate_text(path.read_text(encoding="utf-8", errors="ignore"))
        except Exception as exc:
            print(f"[Extract] Error reading text file {path}: {exc}")
            return None

    if suffix == ".pdf":
        extracted = _extract_text_from_pdf(path)
        if extracted:
            return extracted
    if suffix == ".docx":
        extracted = _extract_text_from_docx(path)
        if extracted:
            return extracted
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        extracted = _extract_text_from_xlsx(path)
        if extracted:
            return extracted
    if suffix in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp"}:
        extracted = _extract_text_from_image(path)
        if extracted:
            return extracted

    if attachment_type and attachment_type.startswith("image/"):
        extracted = _extract_text_from_image(path)
        if extracted:
            return extracted

    if suffix == ".rtf":
        try:
            text = path.read_text(encoding="utf-8", errors="ignore")
            return _truncate_text(text)
        except Exception as exc:
            print(f"[Extract] Error reading RTF {path}: {exc}")
            return None

    print(f"[Extract] Unsupported file type for extraction: {path} ({attachment_type})")
    return None


def _build_attachment_context(attachments: list[dict]) -> tuple[list[dict], str]:
    contexts = []
    lines = []
    deps = _get_extraction_dependencies()
    missing_libs = [name for name, available in deps.items() if not available]
    for attachment in attachments:
        filename = attachment.get("filename") or "unknown"
        content_type = attachment.get("type") or "unknown"
        url = attachment.get("url") or ""
        extracted_text = None
        file_path = _get_uploaded_file_path(url)
        if file_path is not None:
            extracted_text = _extract_text_from_file(file_path, content_type)
        if file_path is None:
            attachment_note = "[Đường dẫn tệp không tìm thấy trên server, nên không thể trích xuất nội dung.]"
        elif extracted_text is None:
            if missing_libs:
                attachment_note = (
                    "[Không thể trích xuất nội dung tự động. Thiếu thư viện hoặc công cụ: "
                    + ", ".join(missing_libs)
                    + ".]"
                )
            else:
                attachment_note = "[Không thể trích xuất nội dung tự động với định dạng hiện tại.]"
        else:
            attachment_note = extracted_text

        contexts.append({
            "filename": filename,
            "url": url,
            "type": content_type,
            "text": extracted_text,
            "path": str(file_path) if file_path is not None else None,
        })
        lines.append(f"Tệp đính kèm: {filename} ({content_type})\n{attachment_note}")

    summary = "\n\n".join(lines)
    if summary:
        summary = "Dưới đây là nội dung trích xuất từ tệp đính kèm:\n\n" + summary
    return contexts, summary


# ══════════════════════════════════════════════════════════════════
#  ROUTES
# ══════════════════════════════════════════════════════════════════
@app.route("/", methods=["GET"])
def index():
    """Serve the main chat UI."""
    if os.path.exists(os.path.join(app.template_folder, "index.html")):
        return render_template("index.html")
    return """
    <!DOCTYPE html>
    <html>
    <head><title>Chatbot</title></head>
    <body>
        <h1>Loading...</h1>
        <p>Please wait for templates to load.</p>
    </body>
    </html>
    """


@app.route("/api/upload", methods=["POST"])
def upload_file():
    try:
        uploaded_file = request.files.get("file")
        if uploaded_file is None:
            return jsonify({"error": "Không có tệp nào được gửi."}), 400
        if uploaded_file.filename == "":
            return jsonify({"error": "Tên tệp không hợp lệ."}), 400

        upload_dir = Path(app.static_folder) / "uploads"
        upload_dir.mkdir(parents=True, exist_ok=True)

        original_filename = secure_filename(uploaded_file.filename)
        extension = Path(original_filename).suffix or ""
        safe_name = f"{uuid.uuid4().hex}{extension}"
        filepath = upload_dir / safe_name
        uploaded_file.save(filepath)

        return jsonify({
            "filename": original_filename,
            "url": f"/static/uploads/{safe_name}",
            "type": uploaded_file.content_type or "application/octet-stream",
            "size": filepath.stat().st_size,
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@app.route("/api/chat", methods=["POST"])
def chat():
    """
    POST endpoint for chat.
    
    Request JSON:
        {
            "message": str,
            "session_id": str (optional),
            "provider": str (optional)
        }
    
    Response JSON:
        {
            "session_id": str,
            "reply": str,
            "intent": str,
            "elapsed": float,
            "provider": str,
            "error": str (if error)
        }
    """
    try:
        data = request.get_json() or {}
        message = data.get("message", "").strip()
        session_id = data.get("session_id") or datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )
        provider = data.get("provider")
        user_id = data.get("user_id")

        if not message and not data.get("attachments"):
            return jsonify({"error": "Message is required"}), 400

        turn_id = int(time.time() * 1000)
        # Ignore attachments for now to keep the chat server running.
        attachment_contexts = []
        message_payload = message

        if history_store.load_session(session_id) is None:
            history_store.create_session(session_id, user_id=user_id)

        result = handle_query(session_id, message_payload, provider_override=provider, user_id=user_id)
        current_provider = _provider or provider or "unknown"

        history_store.append_turn(session_id, {
            "turn_id": turn_id,
            "timestamp": datetime.now().isoformat(),
            "user_query": message,
            "final_answer": result.get("final_answer", ""),
            "intent": result.get("intent", ""),
            "agent_timings": result.get("agent_timings", []),
            "log": result.get("log", []),
            "attachments": attachment_contexts,
        })

        return jsonify({
            "session_id": result["session_id"],
            "reply": result["final_answer"],
            "intent": result["intent"],
            "elapsed": result["elapsed"],
            "provider": current_provider,
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        error_details = getattr(e, "details", None) or getattr(e, "raw_error", None) or getattr(e, "args", [None])[0]
        quota_type = getattr(e, "quota_type", None)
        hint = None
        if quota_type == "requests_per_day":
            hint = "Groq requests per day quota exceeded. Hãy thử lại sau 24h hoặc nâng cấp gói Groq."
        elif quota_type == "tokens_per_day":
            hint = "Groq tokens per day quota exceeded. Hãy thử lại sau 24h hoặc giảm kích thước yêu cầu."
        elif quota_type == "requests_per_minute":
            hint = "Groq requests per minute quota exceeded. Vui lòng đợi một vài giây và thử lại."
        status_code = 429 if quota_type else 500
        return jsonify({
            "error": str(e),
            "error_type": type(e).__name__,
            "error_details": error_details,
            "quota_type": quota_type,
            "hint": hint,
        }), status_code


@app.route("/api/health", methods=["GET"])
def health():
    """Health check endpoint."""
    try:
        if _graph is None:
            init_chatbot()
        available_providers = []
        model_fast = None
        model_main = None
        module = _chatbot_module
        if module is None:
            import importlib.util
            spec = importlib.util.spec_from_file_location(
                "chatbot_module",
                os.path.join(os.path.dirname(__file__), "multi_agent_chatbot_v3.py")
            )
            module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(module)
        if getattr(module, "GROQ_KEY", None):
            available_providers.append("groq")
        if getattr(module, "ANTHROPIC_KEY", None):
            available_providers.append("anthropic")
        model_fast = getattr(module, "MODEL_FAST", None)
        model_main = getattr(module, "MODEL_MAIN", None)
        current_provider = (
            _provider or getattr(module, "PROVIDER", None) or "unknown"
        )
        deps = _get_extraction_dependencies()
        extraction_available = all([
            deps["PIL"],
            deps["tesseract_binary"],
            (deps["pytesseract"] or deps["easyocr"]),
            ENABLE_OCR,
        ])
        return jsonify({
            "status": "ok",
            "provider": current_provider,
            "available_providers": available_providers,
            "model_fast": model_fast,
            "model_main": model_main,
            "extraction_dependencies": deps,
            "tesseract_path": _find_tesseract_binary_path(),
            "extraction_available": extraction_available,
            "ocr_enabled": ENABLE_OCR,
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/debug/groq", methods=["GET"])
def debug_groq():
    """Return active Groq provider state and last retry history."""
    try:
        if _graph is None:
            init_chatbot()

        module = _chatbot_module
        if module is None:
            return jsonify({
                "status": "error",
                "message": "Chatbot module not loaded yet.",
            }), 500

        current_provider = _provider or getattr(module, "PROVIDER", None) or "unknown"
        groq_debug = None
        if hasattr(module, "GroqLLMWrapper"):
            groq_debug = module.GroqLLMWrapper.get_last_debug_state()

        return jsonify({
            "status": "ok",
            "provider": current_provider,
            "use_groq": current_provider == "groq",
            "groq_key_configured": bool(getattr(module, "GROQ_KEY", None)),
            "groq_last_debug": groq_debug or {},
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/sessions", methods=["GET"])
def sessions():
    """List available stored chat sessions."""
    try:
        return jsonify({"sessions": history_store.list_sessions()})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/session/new", methods=["POST"])
def new_session():
    """Create a new chat session ID."""
    try:
        data = request.get_json() or {}
        session_id = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        user_id = data.get("user_id")
        history_store.create_session(session_id, user_id=user_id)
        return jsonify({"session_id": session_id})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route("/api/history/<session_id>", methods=["GET"])
def history(session_id: str):
    """Return saved chat history for a session."""
    try:
        session = history_store.load_session(session_id)
        if session is None:
            return jsonify({"error": "Phiên không tồn tại."}), 404
        return jsonify({
            "session_id": session_id,
            "user_id": session.get("user_id", ""),
            "history": history_store.get_history_text(session_id, 50),
            "turns": session.get("turns", []),
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@app.route("/api/process/<session_id>", methods=["GET"])
def process(session_id: str):
    """Return multi-agent process data for the latest or a specific turn."""
    try:
        turn_id = request.args.get("turn_id")
        if turn_id is not None:
            try:
                turn_id = int(turn_id)
            except ValueError:
                return jsonify({"error": "turn_id phải là số nguyên."}), 400
        report = process_inspector.get_process(session_id, turn_id)
        if report is None:
            return jsonify({"error": "Phiên hoặc lượt không tồn tại."}), 404
        return jsonify(report)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


@app.route("/process/<session_id>", methods=["GET"])
def process_page(session_id: str):
    """Render a simple HTML page showing the multi-agent process."""
    turn_id = request.args.get("turn_id")
    if turn_id is not None:
        try:
            turn_id = int(turn_id)
        except ValueError:
            turn_id = None
    return process_inspector.render_html(session_id, turn_id)


@app.route("/api/users", methods=["GET"])
def list_users():
    try:
        return jsonify({"users": user_store.list_users()})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/users/<user_id>", methods=["GET", "POST"])
def user_profile(user_id: str):
    try:
        if request.method == "GET":
            return jsonify({"user_id": user_id, "profile": user_store.get_user(user_id)})

        payload = request.get_json() or {}
        if not isinstance(payload, dict):
            return jsonify({"error": "Dữ liệu phải là JSON object."}), 400
        user_store.save_user(user_id, payload)
        return jsonify({"user_id": user_id, "profile": user_store.get_user(user_id)})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ══════════════════════════════════════════════════════════════════
#  STARTUP
# ══════════════════════════════════════════════════════════════════
def _validate_provider_keys():
    """Lightweight validation of configured provider API keys.
    This only checks for presence of keys in the chatbot module and prints
    clear guidance to stdout so deployment logs make the issue obvious.
    """
    try:
        import importlib.util
        spec = importlib.util.spec_from_file_location(
            "chatbot_module",
            os.path.join(os.path.dirname(__file__), "multi_agent_chatbot_v3.py"),
        )
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
    except Exception:
        print("[WARN] Could not load chatbot module for key validation.")
        return

    groq_ok = bool(getattr(module, "GROQ_KEY", None))
    anthropic_ok = bool(getattr(module, "ANTHROPIC_KEY", None))
    selected = getattr(module, "PROVIDER", None) or PROVIDER or "unknown"

    if selected in (None, "none", "unknown"):
        if not (groq_ok or anthropic_ok):
            print("[WARN] No LLM provider API key found. Set GROQ_API_KEY or ANTHROPIC_API_KEY in the host environment.")
        else:
            print("[INFO] Provider not explicitly selected; available providers: ", end="")
            avail = []
            if groq_ok:
                avail.append("groq")
            if anthropic_ok:
                avail.append("anthropic")
            print(",".join(avail))
    else:
        if selected == "groq" and not groq_ok:
            print("[WARN] GROQ selected but GROQ_API_KEY not configured or empty on host.")
        if selected == "anthropic" and not anthropic_ok:
            print("[WARN] Anthropic selected but ANTHROPIC_API_KEY not configured or empty on host.")

if __name__ == "__main__":
    print("[INFO] Starting Chatbot Web Server...")
    print("[INFO] Open http://localhost:5000 in your browser")
    # Validate provider keys early so deployment logs show clear guidance
    try:
        _validate_provider_keys()
    except Exception:
        pass
    # Use 0.0.0.0 for Replit/production, 127.0.0.1 for local
    host = os.getenv("FLASK_HOST", "0.0.0.0")
    port = int(os.getenv("PORT", 5000))
    debug = os.getenv("FLASK_DEBUG", "False").lower() == "true"
    app.run(debug=debug, host=host, port=port, use_reloader=False)
